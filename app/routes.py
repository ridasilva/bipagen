from functools import wraps
import os
import unicodedata
import uuid
from datetime import datetime
from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from werkzeug.utils import secure_filename
from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    session,
    abort,
    Response,
)
from app.models import (
    db,
    Colecao,
    Usuario,
    ROLE_ADMIN,
    ROLE_EDITOR,
    ROLE_VIEWER,
    STATUS_PUBLIC,
    STATUS_PRIVATE,
)
from flask_babel import gettext as _
from app.qr_utils import generate_qr, generate_qr_pdf

main = Blueprint("main", __name__)

UPLOAD_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static", "uploads"
)

FIELD_ALIASES = {
    "nome_cepa": ["Strain Name", "Nome da Cepa", "Nombre de la Cepa"],
    "tipo": ["Type", "Tipo", "Tipo"],
    "genero": ["Genus", "Gênero", "Género"],
    "especie": ["Species", "Espécie", "Especie"],
    "cepa_strain": ["Strain", "Cepa", "Cepa"],
    "origem": ["Origin", "Origem", "Origen"],
    "local_coleta": ["Collection Site", "Local de Coleta", "Lugar de Recolección"],
    "data_isolamento": ["Isolation Date", "Data de Isolamento", "Fecha de Aislamiento"],
    "meio_cultivo": ["Culture Medium", "Meio de Cultivo", "Medio de Cultivo"],
    "temperatura_cultivo": [
        "Cultivation Temperature",
        "Temperatura de Cultivo",
        "Temperatura de Cultivo",
    ],
    "metodo_preservacao": [
        "Preservation Method",
        "Método de Preservação",
        "Método de Preservación",
    ],
    "local_armazenamento": [
        "Storage Location",
        "Local de Armazenamento",
        "Ubicación de Almacenamiento",
    ],
    "responsavel": ["Responsible", "Responsável", "Responsable"],
    "data_cadastro": ["Registration Date", "Data de Cadastro", "Fecha de Registro"],
    "observacoes": ["Observations", "Observações", "Observaciones"],
    "publicacoes": ["Publications", "Publicações", "Publicaciones"],
    "status": ["Status", "Status", "Estado"],
}


def _norm(s):
    s = unicodedata.normalize("NFD", s or "")
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip().lower()


TIPOS_VALIDOS = ["Fungo", "Bactéria", "Levedura", "Vírus", "Outros"]
TIPOS_ALIASES = {
    "fungo": "Fungo", "fungos": "Fungo", "hongo": "Fungo", "hongos": "Fungo",
    "bacteria": "Bactéria", "bacterium": "Bactéria",
    "levedura": "Levedura", "levadura": "Levedura", "yeast": "Levedura",
    "virus": "Vírus",
    "outros": "Outros", "outro": "Outros", "otros": "Outros", "otro": "Outros", "other": "Outros",
}


def _cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _parse_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    s = _cell_str(v)[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except ValueError:
        return None


def process_publicacoes():
    valores = request.form.getlist("publicacao_valor")
    existentes = request.form.getlist("publicacao_existente")
    arquivos = request.files.getlist("publicacao_arquivo")
    entries = []
    for i, valor in enumerate(valores):
        valor = (valor or "").strip()
        existente = existentes[i].strip() if i < len(existentes) else ""
        arquivo = arquivos[i] if i < len(arquivos) else None
        if arquivo and arquivo.filename:
            orig = secure_filename(arquivo.filename) or "arquivo"
            ext = os.path.splitext(orig)[1] or ""
            disk = f"pub_{uuid.uuid4().hex}{ext}"
            os.makedirs(UPLOAD_DIR, exist_ok=True)
            arquivo.save(os.path.join(UPLOAD_DIR, disk))
            entries.append({"tipo": "arquivo", "valor": disk, "display": orig})
        elif valor:
            entries.append({"tipo": "link", "valor": valor})
        elif existente:
            if existente.startswith("arquivo:"):
                raw = existente[len("arquivo:"):]
                if "|" in raw:
                    disk, display = raw.split("|", 1)
                else:
                    disk, display = raw, raw
                entries.append({"tipo": "arquivo", "valor": disk, "display": display})
            else:
                entries.append({"tipo": "link", "valor": existente})
    return Colecao.serialize_publicacoes(entries)


@main.route("/lang/<lang>")
def set_language(lang):
    if lang in ("pt", "en", "es"):
        session["lang"] = lang
    return redirect(request.referrer or url_for("main.index"))


def get_current_user():
    user_id = session.get("user_id")
    if user_id is None:
        return None
    return db.session.get(Usuario, user_id)


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        user = get_current_user()
        if user is None:
            flash(_("Please log in to continue."), "warning")
            return redirect(url_for("main.login"))
        return view(*args, **kwargs)

    return wrapped


def role_required(*roles):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            user = get_current_user()
            if user is None:
                flash(_("Please log in to continue."), "warning")
                return redirect(url_for("main.login"))
            if user.role not in roles:
                abort(403)
            return view(*args, **kwargs)

        return wrapped

    return decorator


@main.app_context_processor
def inject_user():
    from flask_babel import get_locale
    return {"current_user": get_current_user(), "get_locale": get_locale}


@main.route("/")
def index():
    return render_template("index.html")


@main.route("/registro", methods=["GET", "POST"])
def registro():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm", "")
        if not username or not email or not password:
            flash(_("Username, email and password are required."), "danger")
        elif password != confirm:
            flash(_("Passwords do not match."), "danger")
        elif Usuario.query.filter_by(username=username).first():
            flash(_("Username already taken."), "danger")
        elif Usuario.query.filter_by(email=email).first():
            flash(_("Email already registered."), "danger")
        else:
            usuario = Usuario(username=username, email=email, role=ROLE_VIEWER)
            usuario.set_password(password)
            db.session.add(usuario)
            db.session.commit()
            flash(_("Account created! Please log in."), "success")
            return redirect(url_for("main.login"))
    return render_template("registro.html")


@main.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        usuario = Usuario.query.filter_by(username=username).first()
        if usuario and usuario.check_password(password):
            session["user_id"] = usuario.id
            flash(_("Welcome, %(name)s!") % {"name": usuario.username}, "success")
            return redirect(url_for("main.index"))
        flash(_("Invalid username or password."), "danger")
    return render_template("login.html")


@main.route("/logout")
def logout():
    session.pop("user_id", None)
    flash(_("You have been logged out."), "info")
    return redirect(url_for("main.index"))


@main.route("/usuarios")
@role_required(ROLE_ADMIN)
def usuarios():
    usuarios = Usuario.query.order_by(Usuario.id.asc()).all()
    return render_template("usuarios.html", usuarios=usuarios)


@main.route("/usuarios/<int:id>/papel", methods=["POST"])
@role_required(ROLE_ADMIN)
def atualizar_papel(id):
    usuario = db.session.get(Usuario, id)
    if not usuario:
        abort(404)
    novo_papel = request.form.get("role", "")
    if novo_papel not in (ROLE_ADMIN, ROLE_EDITOR, ROLE_VIEWER):
        flash(_("Invalid role."), "danger")
        return redirect(url_for("main.usuarios"))
    if usuario.id == get_current_user().id and novo_papel != ROLE_ADMIN:
        flash(_("You cannot remove your own admin role."), "danger")
        return redirect(url_for("main.usuarios"))
    usuario.role = novo_papel
    db.session.commit()
    flash(
        _("Role of %(username)s updated to %(role)s.") % {"username": usuario.username, "role": novo_papel},
        "success",
    )
    return redirect(url_for("main.usuarios"))


@main.route("/usuarios/<int:id>/excluir", methods=["POST"])
@role_required(ROLE_ADMIN)
def excluir_usuario(id):
    usuario = db.session.get(Usuario, id)
    if not usuario:
        abort(404)
    if usuario.id == get_current_user().id:
        flash(_("You cannot delete your own account."), "danger")
        return redirect(url_for("main.usuarios"))
    Colecao.query.filter_by(owner_id=usuario.id).update({Colecao.owner_id: None})
    db.session.delete(usuario)
    db.session.commit()
    flash(_("User %(username)s deleted.") % {"username": usuario.username}, "success")
    return redirect(url_for("main.usuarios"))


def visible_query(user):
    if user is None:
        return Colecao.query.filter(Colecao.status == STATUS_PUBLIC)
    if user.is_admin:
        return Colecao.query
    return Colecao.query.filter(
        db.or_(Colecao.status == STATUS_PUBLIC, Colecao.owner_id == user.id)
    )


def can_view_registro(registro, user):
    if registro.status == STATUS_PUBLIC:
        return True
    if user is None:
        return False
    return user.is_admin or registro.owner_id == user.id


ORDENACAO_COLUNAS = {
    "codigo_unico": Colecao.codigo_unico,
    "nome_cepa": Colecao.nome_cepa,
    "tipo": Colecao.tipo,
    "genero": Colecao.genero,
    "especie": Colecao.especie,
}


@main.route("/colecao")
def listar():
    search = request.args.get("search", "")
    page = request.args.get("page", 1, type=int)
    sort = request.args.get("sort", "codigo_unico")
    order = request.args.get("order", "desc")
    if sort not in ORDENACAO_COLUNAS:
        sort = "codigo_unico"
    if order not in ("asc", "desc"):
        order = "desc"
    per_page = 20
    query = visible_query(get_current_user())
    if search:
        query = query.filter(
            db.or_(
                Colecao.codigo_unico.ilike(f"%{search}%"),
                Colecao.nome_cepa.ilike(f"%{search}%"),
                Colecao.genero.ilike(f"%{search}%"),
                Colecao.especie.ilike(f"%{search}%"),
                Colecao.tipo.ilike(f"%{search}%"),
            )
        )
    coluna = ORDENACAO_COLUNAS[sort]
    sentido = coluna.desc() if order == "desc" else coluna.asc()
    pagination = query.order_by(sentido, Colecao.id.asc()).paginate(
        page=page, per_page=per_page, error_out=False
    )
    registros = pagination.items
    return render_template(
        "listar.html",
        registros=registros,
        pagination=pagination,
        search=search,
        sort=sort,
        order=order,
    )


@main.route("/colecao/novo", methods=["GET", "POST"])
@role_required(ROLE_ADMIN, ROLE_EDITOR)
def novo():
    user = get_current_user()
    if request.method == "POST":
        try:
            registro = Colecao(
                nome_cepa=request.form["nome_cepa"],
                tipo=request.form["tipo"],
                genero=request.form["genero"],
                especie=request.form["especie"],
                cepa_strain=None,
                origem=None,
                local_coleta=None,
                data_isolamento=None,
                responsavel=None,
                data_cadastro=None,
                meio_cultivo=request.form.get("meio_cultivo"),
                temperatura_cultivo=request.form.get("temperatura_cultivo"),
                metodo_preservacao=request.form.get("metodo_preservacao"),
                local_armazenamento=request.form.get("local_armazenamento"),
                observacoes=request.form.get("observacoes"),
                publicacoes=process_publicacoes(),
                status=STATUS_PUBLIC,
                owner_id=user.id,
            )
            db.session.add(registro)
            db.session.commit()
            generate_qr(registro)
            flash(_("Record added successfully!"), "success")
            return redirect(url_for("main.listar"))
        except Exception as e:
            db.session.rollback()
            flash(_("Error: %(message)s") % {"message": e}, "danger")
    return render_template("form.html", registro=None)


@main.route("/colecao/editar/<int:id>", methods=["GET", "POST"])
@role_required(ROLE_ADMIN, ROLE_EDITOR)
def editar(id):
    user = get_current_user()
    registro = Colecao.query.get_or_404(id)
    if not user.is_admin and registro.owner_id != user.id:
        abort(403)
    if request.method == "POST":
        try:
            registro.nome_cepa = request.form["nome_cepa"]
            registro.tipo = request.form["tipo"]
            registro.genero = request.form["genero"]
            registro.especie = request.form["especie"]
            registro.meio_cultivo = request.form.get("meio_cultivo")
            registro.temperatura_cultivo = request.form.get("temperatura_cultivo")
            registro.metodo_preservacao = request.form.get("metodo_preservacao")
            registro.local_armazenamento = request.form.get("local_armazenamento")
            registro.observacoes = request.form.get("observacoes")
            registro.publicacoes = process_publicacoes()
            db.session.commit()
            generate_qr(registro)
            flash(_("Record updated successfully!"), "success")
            return redirect(url_for("main.listar"))
        except Exception as e:
            db.session.rollback()
            flash(_("Error: %(message)s") % {"message": e}, "danger")
    return render_template("form.html", registro=registro)


@main.route("/colecao/importar", methods=["POST"])
@role_required(ROLE_ADMIN, ROLE_EDITOR)
def importar():
    arquivo = request.files.get("arquivo_xlsx")
    status = request.form.get("status", STATUS_PUBLIC)
    if status not in (STATUS_PUBLIC, STATUS_PRIVATE):
        status = STATUS_PUBLIC
    if not arquivo or not arquivo.filename:
        flash(_("Please upload an XLSX file."), "danger")
        return redirect(url_for("main.novo"))
    try:
        wb = load_workbook(arquivo, data_only=True)
    except Exception as e:
        flash(_("Error reading the XLSX file: %(message)s") % {"message": e}, "danger")
        return redirect(url_for("main.novo"))

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows or not any(c is not None for c in rows[0]):
        flash(_("The file is empty."), "danger")
        return redirect(url_for("main.novo"))

    headers = [_cell_str(c) for c in rows[0]]
    alias_map = {}
    for field, aliases in FIELD_ALIASES.items():
        for a in aliases:
            alias_map[_norm(a)] = field
    header_fields = [alias_map.get(_norm(h)) for h in headers]

    user = get_current_user()
    criados = 0
    erros = []
    novos = []
    with db.session.no_autoflush:
        existentes = {
            r[0] for r in db.session.query(Colecao.nome_cepa).all()
        }
    vistos = set()
    for r_idx, row in enumerate(rows[1:], start=2):
        if row is None or all(not _cell_str(c) for c in row):
            continue
        registro = Colecao(owner_id=user.id, status=status)
        extras = []
        obs_direta = None
        for h_idx, valor in enumerate(row):
            campo = header_fields[h_idx] if h_idx < len(header_fields) else None
            valor_txt = _cell_str(valor)
            if campo is None:
                if valor_txt:
                    extras.append(f"{headers[h_idx]}: {valor_txt}")
            elif campo == "status":
                continue
            elif campo == "observacoes":
                if valor_txt:
                    obs_direta = valor_txt
            elif campo == "publicacoes":
                registro.publicacoes = valor_txt or None
            elif campo in ("data_isolamento", "data_cadastro"):
                d = _parse_date(valor)
                if d:
                    setattr(registro, campo, d)
            elif campo == "tipo":
                registro.tipo = (valor_txt[:1].upper() + valor_txt[1:]) if valor_txt else ""
            else:
                setattr(registro, campo, valor_txt or None)
        registro.tipo = registro.tipo or ""
        registro.genero = registro.genero or ""
        registro.especie = registro.especie or ""
        if registro.tipo:
            canonical = TIPOS_ALIASES.get(_norm(registro.tipo))
            if canonical is None:
                erros.append(
                    _("Row %(n)s: unknown type '%(name)s'") % {"n": r_idx, "name": registro.tipo}
                )
                continue
            registro.tipo = canonical
        if not registro.nome_cepa:
            erros.append(_("Row %(n)s: strain name is required") % {"n": r_idx})
            continue
        if registro.nome_cepa in existentes or registro.nome_cepa in vistos:
            erros.append(
                _("Row %(n)s: strain name '%(name)s' already exists")
                % {"n": r_idx, "name": registro.nome_cepa}
            )
            continue
        vistos.add(registro.nome_cepa)
        obs_parts = [p for p in (obs_direta, "; ".join(extras)) if p]
        registro.observacoes = "\n".join(obs_parts) if obs_parts else None
        db.session.add(registro)
        novos.append(registro)
        criados += 1

    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        flash(_("Error importing: some records have duplicate strain names."), "danger")
        return redirect(url_for("main.novo"))

    for registro in novos:
        generate_qr(registro)

    msg = _("%(count)s record(s) imported.") % {"count": criados}
    flash(msg, "success")
    if erros:
        flash(_("%(count)s error(s): ") % {"count": len(erros)} + " | ".join(erros[:8]), "warning")
    return redirect(url_for("main.listar"))


@main.route("/colecao/excluir/<int:id>", methods=["POST"])
@role_required(ROLE_ADMIN)
def excluir(id):
    registro = Colecao.query.get_or_404(id)
    try:
        db.session.delete(registro)
        db.session.commit()
        flash(_("Record deleted successfully!"), "success")
    except Exception as e:
        db.session.rollback()
        flash(_("Error: %(message)s") % {"message": e}, "danger")
    return redirect(url_for("main.listar"))


@main.route("/colecao/excluir-lote", methods=["POST"])
@role_required(ROLE_ADMIN)
def excluir_lote():
    ids = request.form.getlist("ids")
    if not ids:
        flash(_("No records selected."), "warning")
        return redirect(url_for("main.listar"))
    deletados = 0
    try:
        for raw_id in ids:
            try:
                rid = int(raw_id)
            except (TypeError, ValueError):
                continue
            registro = Colecao.query.get(rid)
            if registro:
                db.session.delete(registro)
                deletados += 1
        db.session.commit()
        if deletados:
            flash(_("%(count)s record(s) deleted.") % {"count": deletados}, "success")
    except Exception as e:
        db.session.rollback()
        flash(_("Error: %(message)s") % {"message": e}, "danger")
    return redirect(url_for("main.listar"))


@main.route("/colecao/pdf")
@role_required(ROLE_ADMIN)
def pdf():
    ids = []
    for raw_id in request.args.getlist("ids"):
        try:
            ids.append(int(raw_id))
        except (TypeError, ValueError):
            continue
    query = Colecao.query
    if ids:
        query = query.filter(Colecao.id.in_(ids))
    registros = query.order_by(Colecao.id.asc()).all()
    pdf_buffer = generate_qr_pdf(registros)
    return Response(
        pdf_buffer.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": "attachment; filename=bipagen_qr_codes.pdf"},
    )


@main.route("/colecao/<int:id>")
def detalhe(id):
    registro = Colecao.query.get_or_404(id)
    user = get_current_user()
    if not can_view_registro(registro, user):
        abort(403)
    return render_template("detalhe.html", registro=registro)


@main.route("/servicos", methods=["GET", "POST"])
def servicos():
    categorias = [
        {"id": 1, "nome": _("Category 1 — Basic Strains"), "preco_publico": 200, "preco_privado": 400},
        {"id": 2, "nome": _("Category 2 — Standard Strains"), "preco_publico": 300, "preco_privado": 600},
        {"id": 3, "nome": _("Category 3 — Specialized Strains"), "preco_publico": 500, "preco_privado": 1000},
        {"id": 4, "nome": _("Category 4 — Premium Strains"), "preco_publico": 1050, "preco_privado": 2100},
    ]
    if request.method == "POST":
        ok, message = submit_service_request(categorias)
        flash(message, "success" if ok else "danger")
        return redirect(url_for("main.servicos"))
    return render_template("servicos.html", categorias=categorias)


def submit_service_request(categorias):
    from app.mailer import send_email

    nome = request.form.get("nome", "").strip()
    instituicao = request.form.get("instituicao", "").strip()
    email_req = request.form.get("email", "").strip()
    telefone = request.form.get("telefone", "").strip()
    observacoes = request.form.get("observacoes", "").strip()
    if not nome or not instituicao or not email_req:
        return False, _("Name, institution and email are required.")

    cat_map = {c["id"]: c for c in categorias}
    items = []
    total = 0
    for cid, setor, qty in zip(
        request.form.getlist("categoria"),
        request.form.getlist("setor"),
        request.form.getlist("quantidade"),
    ):
        cat = cat_map.get(int(cid)) if cid else None
        qty = int(qty) if qty else 0
        if not cat or not setor or qty < 1:
            continue
        preco = cat["preco_privado"] if setor == "privado" else cat["preco_publico"]
        subtotal = preco * qty
        total += subtotal
        items.append(
            {
                "nome": cat["nome"],
                "setor": _("Private Sector") if setor == "privado" else _("Public University"),
                "preco": preco,
                "quantidade": qty,
                "subtotal": subtotal,
            }
        )

    if not items:
        return False, _("Add at least one valid requested strain.")

    def brl(v):
        return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    rows = "\n".join(
        "<tr>"
        f"<td>{i['nome']}</td>"
        f"<td>{i['setor']}</td>"
        f"<td>{brl(i['preco'])}</td>"
        f"<td>{i['quantidade']}</td>"
        f"<td>{brl(i['subtotal'])}</td>"
        "</tr>"
        for i in items
    )
    subject = _("BIPAGEN Service Request — %(name)s") % {"name": nome}
    html = f"""
    <h2>{_("BIPAGEN Service Request")}</h2>
    <table cellpadding="6" style="border-collapse:collapse">
      <thead><tr style="border-bottom:2px solid #000">
        <th align="left">{_("Category")}</th><th align="left">{_("Sector")}</th>
        <th align="right">{_("Unit Price")}</th><th align="right">{_("Qty")}</th><th align="right">{_("Subtotal")}</th>
      </tr></thead>
      <tbody>{rows}</tbody>
    </table>
    <p><strong>{_("Total")}: {brl(total)}</strong></p>
    <hr>
    <p><strong>{_("Requester")}:</strong> {nome}</p>
    <p><strong>{_("Institution")}:</strong> {instituicao}</p>
    <p><strong>{_("Email")}:</strong> {email_req}</p>
    <p><strong>{_("Phone")}:</strong> {telefone or '—'}</p>
    <p><strong>{_("Notes")}:</strong> {observacoes or '—'}</p>
    """
    text = (
        _("BIPAGEN Service Request") + "\n"
        + "-----------------------\n"
        + "\n".join(
            f"- {i['nome']} | {i['setor']} | {brl(i['preco'])} x {i['quantidade']} = {brl(i['subtotal'])}"
            for i in items
        )
        + f"\n{_('Total')}: {brl(total)}\n\n"
        + f"{_('Requester')}: {nome}\n{_('Institution')}: {instituicao}\n{_('Email')}: {email_req}\n{_('Phone')}: {telefone or '—'}\n{_('Notes')}: {observacoes or '—'}"
    )
    return send_email(subject, html, text_body=text)
